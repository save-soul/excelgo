#!/usr/bin/env python3
# validate_generated.py — 用 openpyxl 独立交叉校验 Go 库生成的 xlsx。
# openpyxl 是完全独立的实现，从其视角读取即代表 Excel/WPS 打开后的真实结果，
# 用于证明 excelgo 生成的文件与「预期」一致（而非用库自身回读造成的循环验证）。

import os
import zipfile
import sys
import glob
import xml.dom.minidom as minidom

import openpyxl

GEN = os.path.join(os.path.dirname(__file__), "generated")
fails = []


def xml_wellformed(path):
    """用严格 XML 解析器逐个校验压缩包内所有 XML 部件是否良构（Excel/WPS 会拒绝非法 XML）。"""
    z = zipfile.ZipFile(path)
    bad = []
    for n in z.namelist():
        if not n.endswith(".xml") and not n.endswith(".rels"):
            continue
        try:
            minidom.parseString(z.read(n))
        except Exception as e:  # noqa: BLE001
            bad.append(f"{n}: {e}")
    return bad


def check(name, cond, detail=""):
    status = "PASS" if cond else "FAIL"
    print(f"[{status}] {name}" + (f"  -> {detail}" if detail else ""))
    if not cond:
        fails.append(name)


def media_count(path):
    z = zipfile.ZipFile(path)
    return sum(1 for n in z.namelist() if n.startswith("xl/media/"))


# 1) 单元格各类型（独立读取）
p = os.path.join(GEN, "02_cells.xlsx")
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb["Sheet1"]
check("02 A1 共享字符串", ws["A1"].value == "你好", repr(ws["A1"].value))
check("02 A2 整数", ws["A2"].value == 42, repr(ws["A2"].value))
check("02 A3 浮点", abs(float(ws["A3"].value) - 3.14) < 1e-9, repr(ws["A3"].value))
check("02 A4 布尔", ws["A4"].value is True, repr(ws["A4"].value))
# A5 内联字符串：openpyxl 也以普通字符串呈现
check("02 A5 内联字符串", ws["A5"].value == "inline", repr(ws["A5"].value))
# A6 由库写成 t="str" 的公式结果（<v>84</v>），openpyxl 以字符串 "84" 呈现
check("02 A6 公式缓存值", str(ws["A6"].value) == "84", repr(ws["A6"].value))
check("02 A7 自动推断字符串", ws["A7"].value == "auto", repr(ws["A7"].value))

# 2) 区域读写
p = os.path.join(GEN, "03_range.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb["Sheet1"]
got = [[ws.cell(row=r, column=c).value for c in range(2, 5)] for r in range(2, 5)]
exp = [["a", "b", "c"], [1, 2, 3], ["x", "y", "z"]]
check("03 区域 B2:D4", got == exp, repr(got))

# 3) 工作表生命周期（删改后最终仅剩 Base）
p = os.path.join(GEN, "01_lifecycle.xlsx")
wb = openpyxl.load_workbook(p)
check("01 最终工作表=[Base]", wb.sheetnames == ["Base"], repr(wb.sheetnames))

# 4) 行列平移（插入行→删除→插入列 后 A2 应为 mid）
#    用 data_only=True 读取公式缓存值（库不重算，缓存值保留 "topmid"）
p = os.path.join(GEN, "04c_after_insertcol.xlsx")
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb["Sheet1"]
check("04 A1=top", ws["A1"].value == "top", repr(ws["A1"].value))
check("04 A2=mid（数据未被误删）", ws["A2"].value == "mid", repr(ws["A2"].value))
check("04 A3 公式结果", ws["A3"].value == "topmid", repr(ws["A3"].value))

# 5) 图片数量（直接数 zip 内 media 部件，独立于解析器）
p = os.path.join(GEN, "06_pictures.xlsx")
check("06 media 部件 >=2", media_count(p) >= 2, f"count={media_count(p)}")

# 6) 合并工作簿（在 target 基础上 +2 个工作表）
src_target = os.path.join(os.path.dirname(__file__), "testfixtures", "target.xlsx")
base_n = len(openpyxl.load_workbook(src_target).sheetnames)
p = os.path.join(GEN, "08_merge.xlsx")
merged_n = len(openpyxl.load_workbook(p).sheetnames)
check("08 合并后工作表数 = 基线+2", merged_n == base_n + 2, f"{base_n} -> {merged_n}")

# 7) 表格与批注
p = os.path.join(GEN, "10_table_comment.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb["Sheet1"]
check("10 表格存在", len(ws.tables) >= 1, repr(list(ws.tables.keys())))
# 批注：openpyxl 无 ws.comments 属性，直接从 zip 内 xl/comments/* 部件校验
z = zipfile.ZipFile(p)
comment_parts = [n for n in z.namelist() if n.startswith("xl/comments/")]
comment_text_ok = any("备注内容" in z.read(n).decode("utf-8", "replace") for n in comment_parts)
check("10 批注存在", len(comment_parts) >= 1 and comment_text_ok,
      repr(comment_parts))

# 8) 富装饰：合并/冻结/筛选/数据验证/保护/超链接
p = os.path.join(GEN, "11_decor.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb["Sheet1"]
# 测试先 MergeCells 再 UnmergeCells 同一区域，终态应为「无合并单元格」，
# openpyxl 能正常解析即证明取消合并未破坏 XML。
check("11 取消合并后无残留合并", len(ws.merged_cells.ranges) == 0, str(ws.merged_cells.ranges))
check("11 冻结窗格 A2", ws.freeze_panes == "A2", repr(ws.freeze_panes))
check("11 自动筛选", ws.auto_filter.ref is not None, repr(ws.auto_filter.ref))
check("11 数据验证", len(ws.data_validations.dataValidation) >= 1,
      str(len(ws.data_validations.dataValidation)))
check("11 工作表保护", ws.protection.sheet is True, repr(ws.protection.sheet))
# 超链接：openpyxl 3.1.5 无 ws.hyperlinks 属性，直接从 worksheet XML 校验
ws_xml = zipfile.ZipFile(p).read("xl/worksheets/sheet1.xml").decode("utf-8", "replace")
has_hl = "<hyperlink " in ws_xml
check("11 超链接", has_hl, "存在" if has_hl else "worksheet 未见 <hyperlink")

# 9) 文档属性
p = os.path.join(GEN, "13_docprops.xlsx")
wb = openpyxl.load_workbook(p)
check("13 creator=excelgo-test", (wb.properties.creator or "") == "excelgo-test",
      repr(wb.properties.creator))

# 10) 对象式 API 生成文件（14）
p = os.path.join(GEN, "14_object_api.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb["Obj"]
check("14 Obj!A1=对象式", ws["A1"].value == "对象式", repr(ws["A1"].value))
check("14 Obj!A2=123", ws["A2"].value == 123, repr(ws["A2"].value))

# 0) 全部生成文件的 XML 良构性（独立严格解析，任何非法 XML 立即暴露）
for fp in sorted(glob.glob(os.path.join(GEN, "*.xlsx"))):
    bad = xml_wellformed(fp)
    name = os.path.basename(fp)
    check(f"00 良构性 {name}", not bad, ("; ".join(bad) if bad else ""))

print()
if fails:
    print(f"交叉校验失败 {len(fails)} 项：{fails}")
    sys.exit(1)
print("全部交叉校验通过 ✅")
